"""Audit v2 evidence and export paired statistics. No model calls or repairs."""
import csv
import hashlib
import json
import sys
from collections import Counter
from pathlib import Path

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'benchmark'))
from tasks import evaluation_tasks, grade, SKILLS
from architecture import discover_skill

DATA = ROOT / 'results/evaluation-v2'
OUT = ROOT / 'results/analysis-v2'
ORDER = ['single', 'sequential', 'parallel', 'hierarchical', 'reflexive', 'sage',
         'hierarchical_all_skills', 'sage_single', 'sage_all_skills']
LABELS = {
    'es': ['Agente único', 'Secuencial', 'Paralela', 'Jerárquica', 'Reflexiva', 'SAGE-CREH',
           'Jerárquica + 10 skills', 'SAGE sin delegación', 'SAGE + 10 skills'],
    'en': ['Single agent', 'Sequential', 'Parallel', 'Hierarchical', 'Reflexive', 'SAGE-CREH',
           'Hierarchy + 10 skills', 'SAGE without delegation', 'SAGE + 10 skills'],
}


def csv_write(path, rows):
    with path.open('w', newline='') as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def audit(directory):
    manifest = json.loads((directory / 'manifest.json').read_text())
    completed = json.loads((directory / 'completed.json').read_text())
    assert manifest['source_sha256'] == completed['source_sha256'], 'Changed runtime'
    for name, digest in manifest['source_sha256'].items():
        assert hashlib.sha256((directory / 'frozen-source' / name).read_bytes()).hexdigest() == digest
    rows = [json.loads(line) for line in (directory / 'runs.jsonl').read_text().splitlines()]
    assert len(rows) == manifest['run_count'] == len({r['run_id'] for r in rows})
    assert [dict(task_id=r['task_id'], trial=r['trial'], condition=r['condition']) for r in rows] == manifest['order']
    traces = {}
    for row in rows:
        calls = [json.loads(p.read_text()) for p in sorted((directory / 'traces' / row['run_id']).glob('*.json'))]
        assert len(calls) == row['calls']
        for call in calls:
            assert 'Authorization' not in json.dumps(call) and 'Bearer ' not in json.dumps(call)
            request = call['request']
            packet = json.loads(request['input'][1]['content'])
            assert set(packet) <= {'task', 'skill_catalog', 'loaded_skills', 'handoff'}
            assert request['model'] == manifest['model']
            assert request['reasoning']['effort'] == 'low'
            assert request['max_output_tokens'] == 1536
            usage = call.get('response', {}).get('usage')
            if usage:
                assert usage['input_tokens'] + usage['output_tokens'] == usage['total_tokens']
                assert 0 <= usage.get('input_tokens_details', {}).get('cached_tokens', 0) <= usage['input_tokens']
                assert 0 <= usage.get('output_tokens_details', {}).get('reasoning_tokens', 0) <= usage['output_tokens']
        for key in ['input_tokens', 'output_tokens', 'total_tokens']:
            assert sum((c.get('response', {}).get('usage') or {}).get(key, 0) for c in calls) == row['usage'][key]
        traces[row['run_id']] = calls
    return rows, manifest, completed, traces


def choose_development():
    rows, manifest, completed, _ = audit(ROOT / 'results/development-v2')
    stats = []
    for condition in manifest['conditions']:
        group = [row for row in rows if row['condition'] == condition]
        stats.append(dict(condition=condition, runs=len(group), successes=sum(r['success'] for r in group),
                          mean_tokens=np.mean([r['usage']['total_tokens'] for r in group]),
                          mean_seconds=np.mean([r['wall_seconds'] for r in group])))
    chosen = min(stats, key=lambda r: (-r['successes'], r['mean_tokens'], r['mean_seconds']))
    from datetime import datetime, timezone
    decision = dict(decided_at=datetime.now(timezone.utc).isoformat(),
                    rule='most strict successes, then lowest mean tokens, then lowest mean wall seconds',
                    policy=chosen['condition'].split('_')[1], candidates=stats,
                    development_tokens=sum(r['usage']['total_tokens'] for r in rows),
                    development_calls=sum(r['calls'] for r in rows), development_runs=len(rows),
                    development_finished_at=completed['finished_at'],
                    evidence_sha256=hashlib.sha256((ROOT / 'results/development-v2/runs.jsonl').read_bytes()).hexdigest())
    destination = ROOT / 'results/policy-decision-v2.json'
    if destination.exists():
        raise SystemExit('Policy already recorded; do not overwrite the decision.')
    destination.write_text(json.dumps(decision, indent=2))
    print(json.dumps(decision, indent=2))


def main():
    rows, manifest, completed, traces = audit(DATA)
    tasks = {t.id: t for t in evaluation_tasks()}
    assert json.loads((DATA / 'tasks.json').read_text()) == json.loads((ROOT / 'results/evaluation/tasks.json').read_text())
    decision = json.loads((ROOT / 'results/policy-decision-v2.json').read_text())
    assert decision['policy'] == manifest['sage_policy']
    assert decision['decided_at'] < manifest['created_at']
    assert decision['evidence_sha256'] == hashlib.sha256((ROOT / 'results/development-v2/runs.jsonl').read_bytes()).hexdigest()
    sensitivity = []
    transitions = []
    for row in rows:
        task = tasks[row['task_id']]
        if not row['error']:
            assert grade(task, row['answer'])['success'] == row['success']
        eligible = (not row['error'] and task.family == 'sql' and isinstance(row['answer'], dict)
                    and set(row['answer']) == {'answer'} and isinstance(row['answer']['answer'], str))
        row['normalized_success'] = grade(task, {'answer': {'sql': row['answer']['answer']}})['success'] if eligible else row['success']
        row['wrapper_normalized'] = eligible
        sensitivity.append(dict(run_id=row['run_id'], condition=row['condition'], task_id=task.id,
                                strict_success=row['success'], normalized_success=row['normalized_success'], wrapper_normalized=eligible))
        for call in traces[row['run_id']]:
            packet = json.loads(call['request']['input'][1]['content'])
            assert packet['task'] == task.public
            all_skills = row['condition'] in ['hierarchical_all_skills','sage_all_skills']
            relevant = discover_skill(task.public) if row['condition'].startswith('sage') else task.family
            selected = SKILLS if all_skills else {relevant:SKILLS[relevant]}
            assert packet['loaded_skills'] == {name:skill['instructions'] for name,skill in selected.items()}
            if row['condition'].startswith('sage'):
                assert 'skill_catalog' not in packet
            else:
                assert packet['skill_catalog'] == {name:skill['description'] for name,skill in SKILLS.items()}
        if row['condition'].startswith('sage'):
            proposals = []
            for call in traces[row['run_id']]:
                if call['role'] in ['sage_propose', 'sage_independent'] and call.get('output_text'):
                    try:
                        proposal = json.loads(call['output_text'])
                    except json.JSONDecodeError:
                        continue
                    proposals.append(grade(task, proposal.get('candidate'))['success'])
            transitions.append(dict(run_id=row['run_id'], condition=row['condition'], task_id=task.id,
                                    path=row['routing'].get('path', 'error'),
                                    initial_correct=proposals[0] if proposals else False,
                                    any_proposal_correct=any(proposals), final_correct=row['success'],
                                    corrected=bool(proposals and not proposals[0] and row['success']),
                                    regressed=bool(proposals and proposals[0] and not row['success'])))
    grouped = {(task, condition): [r for r in rows if r['task_id'] == task and r['condition'] == condition]
               for task in tasks for condition in ORDER}
    assert all(len(group) == 2 for group in grouped.values())
    draws = np.random.default_rng(20260914).integers(0, len(tasks), size=(10000, len(tasks)))
    summary, boot = [], {}
    for condition in ORDER:
        group = [r for r in rows if r['condition'] == condition]
        per_task = np.array([[np.mean([r['success'] for r in grouped[(task, condition)]]),
                              np.mean([r['usage']['total_tokens'] for r in grouped[(task, condition)]]),
                              np.mean([r['wall_seconds'] for r in grouped[(task, condition)]]),
                              np.mean([r['normalized_success'] for r in grouped[(task, condition)]])]
                             for task in sorted(tasks)])
        boot[condition] = per_task[draws].mean(axis=1)
        interval = np.percentile(boot[condition], [2.5, 97.5], axis=0)
        row = dict(condition=condition, runs=len(group), successes=sum(r['success'] for r in group),
                   normalized_successes=sum(r['normalized_success'] for r in group),
                   success_percent=100*np.mean([r['success'] for r in group]),
                   normalized_success_percent=100*np.mean([r['normalized_success'] for r in group]),
                   mean_tokens=np.mean([r['usage']['total_tokens'] for r in group]),
                   mean_seconds=np.mean([r['wall_seconds'] for r in group]),
                   median_seconds=np.median([r['wall_seconds'] for r in group]),
                   calls=sum(r['calls'] for r in group),
                   errors=sum(r['error'] is not None for r in group),
                   accounting_complete=all(r['accounting_complete'] for r in group),
                   wrapper_normalizations=sum(r['wrapper_normalized'] for r in group))
        for key in ['total_tokens', 'input_tokens', 'output_tokens', 'cached_input_tokens', 'reasoning_tokens']:
            row[key] = sum(r['usage'][key] for r in group)
        for i, name in enumerate(['success', 'tokens', 'seconds', 'normalized_success']):
            multiplier = 100 if i in [0, 3] else 1
            row[name + '_ci_low'], row[name + '_ci_high'] = (interval[:, i]*multiplier).tolist()
        summary.append(row)
    lookup = {r['condition']: r for r in summary}
    comparisons = []
    pairs = [('sage', c) for c in ORDER if c != 'sage'] + [('hierarchical_all_skills', 'hierarchical')]
    for a, b in pairs:
        delta = boot[a] - boot[b]
        metrics = {'success_pp': (lookup[a]['success_percent']-lookup[b]['success_percent'], delta[:, 0]*100),
                   'normalized_success_pp': (lookup[a]['normalized_success_percent']-lookup[b]['normalized_success_percent'], delta[:, 3]*100),
                   'tokens': (lookup[a]['mean_tokens']-lookup[b]['mean_tokens'], delta[:, 1]),
                   'seconds': (lookup[a]['mean_seconds']-lookup[b]['mean_seconds'], delta[:, 2]),
                   'token_change_percent': ((lookup[a]['mean_tokens']/lookup[b]['mean_tokens']-1)*100, (boot[a][:, 1]/boot[b][:, 1]-1)*100)}
        for metric, (point, samples) in metrics.items():
            lo, hi = np.percentile(samples, [2.5, 97.5])
            comparisons.append(dict(a=a, b=b, metric=metric, difference=point, ci_low=lo, ci_high=hi))
    pareto = {}
    for score in ['success_percent', 'normalized_success_percent']:
        pareto[score] = [a['condition'] for a in summary[:6] if not any(
            b[score] >= a[score] and b['mean_tokens'] <= a['mean_tokens'] and b['mean_seconds'] <= a['mean_seconds']
            and (b[score] > a[score] or b['mean_tokens'] < a['mean_tokens'] or b['mean_seconds'] < a['mean_seconds'])
            for b in summary[:6] if b is not a)]
    flat = [{k: r[k] for k in ['run_id', 'task_id', 'family', 'condition', 'trial', 'success', 'normalized_success', 'wall_seconds', 'calls', 'error', 'accounting_complete']} | r['usage'] for r in rows]
    call_rows = []
    for row in rows:
        for call in traces[row['run_id']]:
            packet = json.loads(call['request']['input'][1]['content'])
            usage = call.get('response', {}).get('usage') or {}
            call_rows.append(dict(run_id=row['run_id'],task_id=row['task_id'],condition=row['condition'],
                                  role=call['role'],index=call['index'],seconds=call['seconds'],
                                  loaded_skills='|'.join(packet['loaded_skills']),skill_body_count=len(packet['loaded_skills']),
                                  catalog_count=len(packet.get('skill_catalog',{})),
                                  input_tokens=usage.get('input_tokens'),output_tokens=usage.get('output_tokens'),
                                  total_tokens=usage.get('total_tokens'),
                                  cached_input_tokens=usage.get('input_tokens_details',{}).get('cached_tokens'),
                                  reasoning_tokens=usage.get('output_tokens_details',{}).get('reasoning_tokens'),
                                  error=call.get('error')))
    by_task = [dict(task_id=t, family=tasks[t].family, condition=c,
                    strict=sum(r['success'] for r in group), normalized=sum(r['normalized_success'] for r in group),
                    mean_tokens=np.mean([r['usage']['total_tokens'] for r in group]), mean_seconds=np.mean([r['wall_seconds'] for r in group]))
               for (t,c),group in grouped.items()]
    result = dict(version='2.0', summary=summary, comparisons=comparisons, pareto=pareto,
                  run_count=len(rows), distinct_tasks=len(tasks), calls=sum(r['calls'] for r in rows),
                  total_tokens=sum(r['usage']['total_tokens'] for r in rows),
                  errors=sum(r['error'] is not None for r in rows), accounting_complete=all(r['accounting_complete'] for r in rows),
                  start=manifest['created_at'], end=completed['finished_at'], bootstrap_draws=10000,
                  policy_decision=decision, source_hashes_verified=True,
                  call_models=dict(Counter(c.get('response', {}).get('model') for calls in traces.values() for c in calls)),
                  routing_counts=dict(Counter(r['condition'] + ':' + r['path'] for r in transitions)),
                  role_counts=dict(Counter(r['condition'] + ':' + c['role'] for r in rows for c in traces[r['run_id']])))
    OUT.mkdir(exist_ok=True)
    (OUT / 'summary.json').write_text(json.dumps(result, indent=2))
    exports = {'summary':summary, 'runs':flat, 'by-task':by_task, 'paired-comparisons':comparisons,
               'format-sensitivity':sensitivity, 'routing':transitions, 'calls':call_rows}
    book = Workbook()
    readme = book.active
    readme.title = 'Read me'
    for row in [
        ['Study / estudio', 'CREH Labs, v2, SAGE-CREH'],
        ['Design / diseño', '12 known tasks x 2 repetitions x 9 arms = 216; NOT a blind holdout'],
        ['Model / modelo', manifest['model'] + ' / low / max output 1536 per call'],
        ['Primary / primaria', 'Strict exact completion; total observed input + output; wall seconds excluding grading'],
        ['Secondary / secundaria', 'Normalized SQL wraps raw answer string only; same query and evaluator'],
        ['Intervals / intervalos', '10000 paired task-cluster percentile draws; both repetitions retained; seed 20260914'],
        ['Tokens', 'Cached input and reasoning output already included; never added twice'],
        ['Development / desarrollo', 'Separate 20 runs; policy chosen before replication'],
        ['Ablations / ablaciones', 'SAGE without delegation; SAGE all ten skills; original hierarchy all ten skills'],
        ['Interpretation / interpretación', 'Prompt and catalog changes confound SAGE versus original topologies; use ablations'],
    ]:
        readme.append(row)
    for name, data in exports.items():
        csv_write(OUT / (name + '.csv'), data)
        sheet = book.create_sheet(name[:31])
        sheet.append(list(data[0]))
        for row in data:
            sheet.append(list(row.values()))
    for sheet in book:
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='405E59')
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = min(50, max(14, max(len(str(c.value or '')) for c in col)+2))
    book.save(OUT / 'benchmark-data.xlsx')
    figures(summary, by_task)
    print(json.dumps({k:v for k,v in result.items() if k not in ['comparisons', 'role_counts']}, indent=2))


def figures(summary, by_task):
    plt.rcParams.update({'font.family':'DejaVu Sans', 'font.size':13, 'text.color':'#262b2d', 'axes.labelcolor':'#262b2d', 'svg.fonttype':'none'})
    for lang in ['es', 'en']:
        fig, axes = plt.subplots(1, 3, figsize=(12.4,4.5), gridspec_kw={'width_ratios':[1.1,1,1]})
        for ax, metric, label in zip(axes, ['success_percent', 'mean_tokens', 'mean_seconds'],
                                   ['Cumplimiento (%)', 'Tokens por tarea', 'Segundos por tarea'] if lang=='es' else ['Task completion (%)','Tokens per task','Seconds per task']):
            for index, row in enumerate(summary[:6]):
                ax.barh(index, row[metric], height=.55, color='#405e59' if index==5 else '#b3c4bd', zorder=3)
                ci = {'success_percent':'success','mean_tokens':'tokens','mean_seconds':'seconds'}[metric]
                ax.errorbar(row[metric], index, xerr=[[max(0,row[metric]-row[ci+'_ci_low'])],[max(0,row[ci+'_ci_high']-row[metric])]], fmt='none', ecolor='#262b2d', capsize=3, zorder=4)
                if metric=='success_percent':
                    ax.scatter(row['normalized_success_percent'], index, marker='D',s=35,facecolors='white',edgecolors='#262b2d',zorder=6)
                text = f"{row[metric]:.1f}" if metric!='mean_tokens' else f"{row[metric]:,.0f}"
                ax.text(max(row[metric]*.06, .01), index, text, va='center', fontsize=12, color='white' if index==5 else '#262b2d', zorder=5)
            ax.set_yticks(range(6), LABELS[lang][:6] if ax is axes[0] else ['']*6)
            ax.invert_yaxis()
            ax.set_xlabel(label)
            ax.grid(axis='x', color='#e7ede9',zorder=0)
            ax.tick_params(axis='y', length=0)
            for spine in ax.spines.values(): spine.set_visible(False)
            if metric=='success_percent': ax.set_xlim(0,108)
        note = ('24 ejecuciones por condición · barras: estricto · rombo: SQL normalizado · intervalos exploratorios 95%' if lang=='es' else '24 runs per condition · bars: strict · diamond: normalized SQL · exploratory 95% intervals')
        fig.text(.5,.01,note,ha='center',fontsize=9,color='#6f6f69')
        fig.tight_layout(rect=(0,.06,1,1))
        for suffix in ['png','pdf','svg']: fig.savefig(OUT/f'comparison-{lang}.{suffix}',dpi=190,facecolor='white')
        plt.close(fig)
        fig, ax = plt.subplots(figsize=(12.4,4.8))
        ax.axis('off')
        headings = ['Configuración','Estricto','SQL normalizado','Tokens/tarea','s/tarea','Llamadas'] if lang=='es' else ['Configuration','Strict','Normalized SQL','Tokens/task','s/task','Calls']
        cells = [[LABELS[lang][i],f"{r['successes']}/24",f"{r['normalized_successes']}/24",f"{r['mean_tokens']:,.0f}",f"{r['mean_seconds']:.2f}",str(r['calls'])] for i,r in enumerate(summary)]
        table = ax.table(cellText=cells,colLabels=headings,cellLoc='center',colWidths=[.31,.12,.18,.15,.12,.12],bbox=[0,0,1,1])
        table.auto_set_font_size(False);table.set_fontsize(10)
        for (r,c),cell in table.get_celld().items():
            cell.set_edgecolor('#cdd5d1');cell.set_linewidth(.6)
            cell.set_facecolor('#405e59' if r==0 else ('#e7ede9' if r==6 else '#ffffff'))
            if r==0: cell.set_text_props(color='white',weight='bold')
            if c==0 and r>0: cell.set_text_props(ha='left')
        fig.tight_layout()
        for suffix in ['png','pdf','svg']:fig.savefig(OUT/f'all-conditions-{lang}.{suffix}',dpi=190,facecolor='white')
        plt.close(fig)


if __name__ == '__main__':
    if '--choose-development' in sys.argv: choose_development()
    else: main()
