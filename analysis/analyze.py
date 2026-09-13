"""Audit recorded evidence and produce tables, paired intervals and figures.

This script never calls a model and never changes benchmark definitions.
"""
from __future__ import annotations

import csv
import hashlib
import json
import sys
from collections import Counter
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "benchmark"))
from tasks import evaluation_tasks, grade
EVALUATION = ROOT / "results" / "evaluation"
OUTPUT = ROOT / "results" / "analysis"
ORDER = ["single", "sequential", "parallel", "hierarchical", "reflexive", "hierarchical_all_skills"]
LABELS = {
    "en": ["Single agent", "Sequential", "Parallel", "Hierarchical", "Reflexive", "Hierarchical\n(all skills)"],
    "es": ["Agente único", "Secuencial", "Paralela", "Jerárquica", "Reflexiva", "Jerárquica\n(todos los skills)"],
}
FAMILIES = {
    "en": ["Ledger", "Calendar", "Dependencies", "Selection", "Memory", "SQL"],
    "es": ["Registros", "Calendario", "Dependencias", "Selección", "Memoria", "SQL"],
}


def csv_write(path, rows):
    with path.open("w", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def main():
    if not (EVALUATION / "completed.json").exists():
        raise SystemExit("Evaluation must finish before final analysis.")
    manifest = json.loads((EVALUATION / "manifest.json").read_text())
    completed = json.loads((EVALUATION / "completed.json").read_text())
    rows = [json.loads(line) for line in (EVALUATION / "runs.jsonl").read_text().splitlines()]
    if len(rows) != manifest["run_count"] or len({r["run_id"] for r in rows}) != len(rows):
        raise SystemExit("Incomplete or duplicate runs.")
    if manifest["source_sha256"] != completed["source_sha256"]:
        raise SystemExit("Benchmark source changed during evaluation.")
    # Protocol prose was amended after evaluation; never rewrite historical hashes.
    amendments = json.loads((ROOT / "editorial-amendments.json").read_text())["files"]
    if set(amendments) != {"protocol.en.md", "protocol.es.md"}:
        raise SystemExit("Only the two protocol documents may have editorial amendments.")
    for filename, digest in manifest["source_sha256"].items():
        current = hashlib.sha256((ROOT / filename).read_bytes()).hexdigest()
        if filename in amendments:
            expected = amendments[filename]
            if digest != expected["original_sha256"] or current != expected["current_sha256"]:
                raise SystemExit("Protocol editorial amendment mismatch: " + filename)
        elif current != digest:
            raise SystemExit("Frozen benchmark source was modified: " + filename)
    # Match the published schedule, rather than silently dropping a bad block.
    actual_order = [dict(task_id=r["task_id"], trial=r["trial"], condition=r["condition"]) for r in rows]
    if actual_order != manifest["order"]:
        raise SystemExit("Recorded order differs from the frozen schedule.")
    task_lookup = {task.id: task for task in evaluation_tasks()}
    # Explicitly post-hoc sensitivity, not a change to the frozen primary grader:
    # accept a raw SQL string under `answer` by wrapping it in {sql: ...}.
    # The SQL text itself is never repaired, reformatted or sent to a model.
    sensitivity_rows = []
    for row in rows:
        envelope = row["answer"]
        eligible = (row["error"] is None and row["family"] == "sql"
                    and isinstance(envelope, dict) and set(envelope) == {"answer"}
                    and isinstance(envelope["answer"], str))
        result = grade(task_lookup[row["task_id"]], {"answer": {"sql": envelope["answer"]}}) if eligible else None
        row["normalized_success"] = result["success"] if result else row["success"]
        row["sql_wrapper_normalized"] = eligible
        sensitivity_rows.append(dict(run_id=row["run_id"], condition=row["condition"], task_id=row["task_id"],
                                     strict_success=row["success"], wrapper_normalized=eligible,
                                     normalized_success=row["normalized_success"],
                                     fixture_passes=json.dumps(result.get("fixture_passes")) if result else ""))
    by_task_condition, call_models, tiers, role_counts = {}, Counter(), Counter(), Counter()
    for row in rows:
        by_task_condition.setdefault((row["task_id"], row["condition"]), []).append(row)
        traces = [json.loads(p.read_text()) for p in sorted((EVALUATION / "traces" / row["run_id"]).glob("*.json"))]
        if len(traces) != row["calls"]:
            raise SystemExit("Trace count does not match run: " + row["run_id"])
        token_sum = 0
        for trace in traces:
            role_counts[row["condition"] + ":" + trace["role"]] += 1
            response = trace.get("response", {})
            usage = response.get("usage") or {}
            token_sum += usage.get("total_tokens", 0)
            if usage and usage["input_tokens"] + usage["output_tokens"] != usage["total_tokens"]:
                raise SystemExit("Provider token arithmetic mismatch.")
            call_models[response.get("model", "unknown")] += 1
            tiers[str(response.get("service_tier"))] += 1
            if "Authorization" in json.dumps(trace) or "Bearer " in json.dumps(trace):
                raise SystemExit("Unexpected authorization material in public trace.")
        if token_sum != row["usage"]["total_tokens"]:
            raise SystemExit("Per-call and per-run usage mismatch.")
    task_ids = sorted({r["task_id"] for r in rows})
    for task_id in task_ids:
        for condition in ORDER:
            if len(by_task_condition.get((task_id, condition), [])) != manifest["repetitions"]:
                raise SystemExit("Unbalanced task block.")
    complete_accounting = all(row["accounting_complete"] for row in rows)
    OUTPUT.mkdir(exist_ok=True)
    rng = np.random.default_rng(20260913)
    draws = rng.integers(0, len(task_ids), size=(10000, len(task_ids)))
    summary, boot = [], {}
    for condition in ORDER:
        group = [row for row in rows if row["condition"] == condition]
        per_task = np.array([[np.mean([r["success"] for r in by_task_condition[(task, condition)]]),
                              np.mean([r["usage"]["total_tokens"] for r in by_task_condition[(task, condition)]]),
                              np.mean([r["wall_seconds"] for r in by_task_condition[(task, condition)]]),
                              np.mean([r["normalized_success"] for r in by_task_condition[(task, condition)]])]
                             for task in task_ids])
        boot[condition] = per_task[draws].mean(axis=1)
        intervals = np.percentile(boot[condition], [2.5, 97.5], axis=0)
        successes = sum(row["success"] for row in group)
        total_tokens = sum(row["usage"]["total_tokens"] for row in group)
        wall = sum(row["wall_seconds"] for row in group)
        summary.append(dict(
            condition=condition, runs=len(group), successes=successes, success_percent=100*successes/len(group),
            success_ci_low=100*intervals[0, 0], success_ci_high=100*intervals[1, 0],
            total_tokens=total_tokens, mean_tokens=total_tokens/len(group),
            tokens_ci_low=intervals[0, 1], tokens_ci_high=intervals[1, 1],
            input_tokens=sum(r["usage"]["input_tokens"] for r in group),
            output_tokens=sum(r["usage"]["output_tokens"] for r in group),
            cached_input_tokens=sum(r["usage"]["cached_input_tokens"] for r in group),
            reasoning_tokens=sum(r["usage"]["reasoning_tokens"] for r in group),
            mean_seconds=wall/len(group), median_seconds=float(np.median([r["wall_seconds"] for r in group])),
            seconds_ci_low=intervals[0, 2], seconds_ci_high=intervals[1, 2], total_wall_seconds=wall,
            calls=sum(r["calls"] for r in group), successes_per_100k_tokens=successes*100000/total_tokens if complete_accounting else None,
            successes_per_minute=successes*60/wall,
            mean_peak_input_tokens=float(np.mean([r["max_call_input_tokens"] for r in group])),
            accounting_complete=all(r["accounting_complete"] for r in group),
            infrastructure_or_protocol_errors=sum(r["error"] is not None for r in group),
            normalized_successes=sum(r["normalized_success"] for r in group),
            normalized_success_percent=100*np.mean([r["normalized_success"] for r in group]),
            normalized_success_ci_low=100*intervals[0, 3], normalized_success_ci_high=100*intervals[1, 3],
            sql_wrapper_normalizations=sum(r["sql_wrapper_normalized"] for r in group),
        ))
    selected, all_skills = [next(row for row in summary if row["condition"] == c) for c in ["hierarchical", "hierarchical_all_skills"]]
    paired = boot["hierarchical_all_skills"] - boot["hierarchical"]
    percent_tokens = (boot["hierarchical_all_skills"][:, 1] / boot["hierarchical"][:, 1] - 1)*100
    ablation = {
        "definition": "all skills minus selected skill, paired by task with both repetitions retained",
        "success_percentage_point_difference": all_skills["success_percent"]-selected["success_percent"],
        "success_difference_ci": list(np.percentile(paired[:, 0]*100, [2.5, 97.5])),
        "mean_token_difference": all_skills["mean_tokens"]-selected["mean_tokens"],
        "token_difference_ci": list(np.percentile(paired[:, 1], [2.5, 97.5])),
        "token_increase_percent": (all_skills["mean_tokens"]/selected["mean_tokens"]-1)*100,
        "token_increase_percent_ci": list(np.percentile(percent_tokens, [2.5, 97.5])),
        "mean_second_difference": all_skills["mean_seconds"]-selected["mean_seconds"],
        "seconds_difference_ci": list(np.percentile(paired[:, 2], [2.5, 97.5])),
        "normalized_success_percentage_point_difference": all_skills["normalized_success_percent"]-selected["normalized_success_percent"],
        "normalized_success_difference_ci": list(np.percentile(paired[:, 3]*100, [2.5, 97.5])),
    }
    pareto = []
    for candidate in summary[:5]:
        dominated = any(other["success_percent"] >= candidate["success_percent"]
                        and other["mean_tokens"] <= candidate["mean_tokens"]
                        and other["mean_seconds"] <= candidate["mean_seconds"]
                        and (other["success_percent"] > candidate["success_percent"]
                             or other["mean_tokens"] < candidate["mean_tokens"]
                             or other["mean_seconds"] < candidate["mean_seconds"])
                        for other in summary[:5] if other is not candidate)
        if not dominated:
            pareto.append(candidate["condition"])
    table_tasks = []
    for task_id in task_ids:
        table_tasks.append(dict(task_id=task_id, family=by_task_condition[(task_id, "single")][0]["family"],
                               **{c: sum(r["success"] for r in by_task_condition[(task_id, c)]) for c in ORDER}))
    flat_rows = []
    for row in rows:
        flat_rows.append({key: row[key] for key in ["run_id", "task_id", "family", "condition", "trial", "success", "score", "wall_seconds", "api_seconds_sum", "calls", "accounting_complete", "reason", "started_at", "finished_at"]}
                         | row["usage"])
    results = dict(summary=summary, ablation=ablation, pareto_primary=pareto, call_models=dict(call_models), service_tiers=dict(tiers),
                   run_count=len(rows), distinct_tasks=len(task_ids), calls=sum(r["calls"] for r in rows),
                   total_tokens=sum(r["usage"]["total_tokens"] for r in rows), accounting_complete=complete_accounting,
                   errors=sum(r["error"] is not None for r in rows), bootstrap_draws=10000,
                   evaluation_start=manifest["created_at"], evaluation_end=completed["finished_at"],
                   source_hashes_verified=True,
                   source_hash_verification="Benchmark files match historical hashes; protocol documents match documented editorial amendment hashes.",
                   role_counts=dict(role_counts),
                   sensitivity_note="Post-hoc: wrap a raw SQL answer string as {sql:string}; no SQL text repair, no rerun, primary grades unchanged.",
                   sql_wrapper_normalizations=sum(r["sql_wrapper_normalized"] for r in rows),
                   normalized_successes=sum(r["normalized_success"] for r in rows))
    (OUTPUT / "summary.json").write_text(json.dumps(results, indent=2))
    csv_write(OUTPUT / "summary.csv", summary)
    csv_write(OUTPUT / "runs.csv", flat_rows)
    csv_write(OUTPUT / "by-task.csv", table_tasks)
    csv_write(OUTPUT / "format-sensitivity.csv", sensitivity_rows)
    workbook(summary, flat_rows, table_tasks, ablation, sensitivity_rows)
    figures(summary, table_tasks)
    print(json.dumps(results, indent=2))


def workbook(summary, rows, by_task, ablation, sensitivity):
    book = Workbook()
    readme = book.active
    readme.title = "Read me"
    for row in [
        ["Subagent orchestration microbenchmark", "Version 1.0 / 2026-09-13"],
        ["Scope", "12 distinct synthetic tasks; 2 repetitions; 6 conditions; live Luna low calls."],
        ["Main comparison", "single, sequential, parallel, hierarchical, reflexive"],
        ["Additional condition", "hierarchical_all_skills is a paired payload ablation."],
        ["Tokens", "Provider-reported input + output; cached input and reasoning are already included."],
        ["Intervals", "95% percentile cluster bootstrap; 10,000 draws; seed 20260913; cluster = task id."],
        ["Timing", "Wall time includes orchestration/network, excludes final grading; not saturated throughput."],
        ["Caution", "Exploratory small-sample results, not a general architecture ranking or production test."],
        ["Post-hoc sensitivity", "Raw SQL strings in answer are wrapped as {sql:string}; original primary scores stay unchanged."],
        ["Español", "12 tareas distintas, dos repeticiones, seis configuraciones. Datos exploratorios."],
        ["Tokens (ES)", "Entrada más salida informadas por la API; razonamiento y caché ya incluidos."],
        ["Tiempo (ES)", "Tiempo de pared de la ejecución, excluyendo el evaluador final."],
    ]:
        readme.append(row)
    for name, data in [("Summary", summary), ("Runs", rows), ("By task", by_task), ("Format sensitivity", sensitivity)]:
        sheet = book.create_sheet(name)
        sheet.append(list(data[0]))
        for row in data:
            sheet.append(list(row.values()))
    sheet = book.create_sheet("Paired ablation")
    sheet.append(["Metric", "Value"])
    for key, value in ablation.items():
        sheet.append([key, json.dumps(value) if isinstance(value, list) else value])
    sources = json.loads((ROOT / "sources.json").read_text())
    sheet = book.create_sheet("Sources")
    sheet.append(["ID", "Title", "Authors", "First submitted", "Version used", "URL", "Scope"])
    for source in sources:
        sheet.append([source["id"], source["title"], source["authors"], source["first_submitted"], source["version"], source["url"], source["use_en"]])
    for sheet in book:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="405E59")
        for column in sheet.columns:
            letter = column[0].column_letter
            width = min(70, max(16, max(len(str(cell.value or "")) for cell in column) + 2))
            sheet.column_dimensions[letter].width = width
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if isinstance(cell.value, float):
                    cell.number_format = "0.00"
    book.save(OUTPUT / "benchmark-data.xlsx")


def figures(summary, task_rows):
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 10,
                         "text.color": "#262b2d", "axes.labelcolor": "#262b2d",
                         "axes.edgecolor": "#cdd5d1", "xtick.color": "#555b58", "ytick.color": "#262b2d",
                         "svg.fonttype": "none", "pdf.fonttype": 42})
    for language in ["en", "es"]:
        figure, axes = plt.subplots(1, 3, figsize=(8.5, 4.1), sharey=True)
        y = np.arange(6)
        for axis, key, low, high, title in zip(
            axes, ["success_percent", "mean_tokens", "mean_seconds"],
            ["success_ci_low", "tokens_ci_low", "seconds_ci_low"],
            ["success_ci_high", "tokens_ci_high", "seconds_ci_high"],
            (["Task success (%)", "Tokens per run", "Seconds per run"] if language == "en" else
             ["Tareas resueltas (%)", "Tokens por ejecución", "Segundos por ejecución"]),
        ):
            values = np.array([r[key] for r in summary])
            errors = np.array([values - [r[low] for r in summary], [r[high] for r in summary] - values])
            axis.barh(y, values, height=.55, color=["#405e59"]*5+["#a9b8b1"])
            axis.errorbar(values, y, xerr=np.maximum(0, errors), fmt="none", ecolor="#262b2d", capsize=2, linewidth=1)
            axis.set_title(title, fontsize=10, loc="left", pad=12)
            axis.set_axisbelow(True)
            axis.xaxis.grid(True, color="#e7ede9", linewidth=.6)
            axis.spines[["top", "right", "left"]].set_visible(False)
            axis.tick_params(axis="y", length=0)
            if key == "success_percent":
                axis.set_xlim(0, 108)
                axis.set_xticks([0, 50, 100])
            else:
                axis.set_xlim(left=0)
                axis.ticklabel_format(axis="x", style="plain")
        axes[0].set_yticks(y, LABELS[language])
        axes[0].invert_yaxis()
        figure.subplots_adjust(left=.23, right=.98, bottom=.18, top=.88, wspace=.3)
        note = ("24 runs per condition · 12 tasks × 2 repetitions\nBars: means; whiskers: task-cluster 95% intervals"
                if language == "en" else "24 ejecuciones por configuración · 12 tareas × 2 repeticiones\nBarras: medias; líneas: intervalos del 95% por tarea")
        figure.text(.02, .025, note, fontsize=8, color="#6f6f69")
        for ext in ["pdf", "png", "svg"]:
            figure.savefig(OUTPUT / f"comparison-{language}.{ext}", dpi=190, facecolor="white")
        plt.close(figure)

        figure, axis = plt.subplots(figsize=(8.5, 4.4))
        matrix = np.array([[sum(row[c] for row in task_rows[2*i:2*i+2]) for c in ORDER] for i in range(6)])
        from matplotlib.colors import ListedColormap, BoundaryNorm
        cmap = ListedColormap(["#f4f0ed", "#e2e7e2", "#bdccc4", "#7f9b8d", "#405e59"])
        axis.imshow(matrix, cmap=cmap, norm=BoundaryNorm(np.arange(-.5, 5.5), 5), aspect="auto")
        for i in range(6):
            for j in range(6):
                axis.text(j, i, f"{matrix[i,j]}/4", ha="center", va="center", color="white" if matrix[i,j]>=3 else "#262b2d")
        axis.set_yticks(range(6), FAMILIES[language])
        axis.set_xticks(range(6), LABELS[language], fontsize=8)
        axis.tick_params(length=0, pad=8)
        axis.spines[:].set_visible(False)
        figure.subplots_adjust(left=.18, right=.98, top=.93, bottom=.2)
        note = "Exact successes by family: two instances × two repetitions" if language == "en" else "Aciertos exactos por grupo: dos instancias × dos repeticiones"
        figure.text(.18, .02, note, fontsize=8, color="#6f6f69")
        for ext in ["pdf", "png", "svg"]:
            figure.savefig(OUTPUT / f"families-{language}.{ext}", dpi=190, facecolor="white")
        plt.close(figure)


if __name__ == "__main__":
    main()
